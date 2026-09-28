from math import cos, pi, sin
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_p2p_documents as base


ROOT = Path(__file__).resolve().parents[1]
EXPORTS = ROOT / "exports"
HEX_BAND = Path("/private/tmp/kwh-partner-brief-hex-band.png")

FONT = "Arial"
INK = "17140F"
BODY = "4A4439"
MUTED = "7F7769"
PAPER = "FCFBF6"
CREAM = "F8F4E8"
LINE = "E6DFCF"
HONEY = "D9A91E"
HONEY_D = "8D6300"
HONEY_P = "F6E8AE"
MOSS = "71823F"
MOSS_P = "EDF1E2"
DARK = "403A2C"
WHITE = "FFFFFF"

CONTENT_DXA = 9360


def rgb(value):
    value = value.replace("#", "")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_run(run, size=10.5, bold=False, italic=False, color=INK, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def configure(doc):
    base.configure_document(
        doc,
        "kWh Electric — Battery OEM Partner Brief v2",
        "A partner-facing introduction to kWh Electric's battery participation infrastructure",
        confidential=False,
    )
    props = doc.core_properties
    props.comments = "External partner brief. Commercial terms are working structures, not quoted prices."

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Heading 1", 15, INK, 12, 6),
        ("Heading 2", 12.5, HONEY_D, 9, 4),
        ("Heading 3", 11.5, MOSS, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        for p in section.header.paragraphs:
            for run in p.runs:
                set_run(run, size=8.5, bold=run.bold, color=MUTED)
        for p in section.footer.paragraphs:
            for run in p.runs:
                if run.text:
                    set_run(run, size=8, bold=False, color=MUTED)


def make_hex_band():
    width, height = 1800, 118
    img = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    radius = 31
    dx = radius * 1.72
    dy = radius * 1.49
    line = (224, 217, 202, 150)
    pale = (246, 232, 174, 110)
    for row in range(3):
        cy = 7 + row * dy
        offset = dx / 2 if row % 2 else 0
        col = -1
        while True:
            cx = offset + col * dx
            if cx - radius > width:
                break
            points = [
                (cx + radius * cos(pi / 3 * i), cy + radius * sin(pi / 3 * i))
                for i in range(6)
            ]
            fill = pale if (row, col) in {(0, 1), (1, 22)} else None
            draw.polygon(points, outline=line, fill=fill, width=2)
            col += 1
    img.save(HEX_BAND)


def add_hex_band(doc, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run()
    shape = run.add_picture(str(HEX_BAND), width=Inches(6.5))
    shape._inline.docPr.set("descr", "Decorative honeycomb pattern")
    shape._inline.docPr.set("title", "Honeycomb pattern")
    return p


def add_kicker(doc, text, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run("⬢  "), 9, bold=True, color=HONEY)
    set_run(p.add_run(text.upper()), 8.5, bold=True, color=HONEY_D)
    return p


def add_title(doc, text, size=27, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size, bold=True, color=INK)
    return p


def add_para(doc, text="", size=10.5, color=BODY, bold=False, italic=False, after=5, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_together = keep
    if align is not None:
        p.alignment = align
    set_run(p.add_run(text), size, bold=bold, italic=italic, color=color)
    return p


def add_rich(doc, parts, size=10.5, after=5, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_together = keep
    if align is not None:
        p.alignment = align
    for text, kwargs in parts:
        set_run(
            p.add_run(text),
            kwargs.get("size", size),
            bold=kwargs.get("bold", False),
            italic=kwargs.get("italic", False),
            color=kwargs.get("color", BODY),
        )
    return p


def set_cell(cell, text, size=9.4, bold=False, color=BODY, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    if fill:
        base.set_cell_shading(cell, fill)
    base.set_cell_border(cell, LINE, size=5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = align
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    set_run(p.add_run(text), size, bold=bold, color=color)
    return p


def add_compare(doc, headers, bodies, widths=(4680, 4680)):
    table = doc.add_table(rows=2, cols=2)
    base.set_table_geometry(table, list(widths))
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, size=10.2, bold=True, color=INK, fill=HONEY_P)
    base.set_repeat_header(table.rows[0])
    for i, body in enumerate(bodies):
        set_cell(table.rows[1].cells[i], body, size=9.5, color=BODY, fill=PAPER)
    add_para(doc, "", after=2)
    return table


def add_metric_strip(doc, metrics):
    widths = [CONTENT_DXA // len(metrics)] * len(metrics)
    widths[-1] += CONTENT_DXA - sum(widths)
    table = doc.add_table(rows=1, cols=len(metrics))
    base.set_table_geometry(table, widths)
    for i, (value, label, fill, accent) in enumerate(metrics):
        cell = table.cell(0, i)
        base.set_cell_shading(cell, fill)
        base.set_cell_border(cell, LINE, size=5)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(value), 15, bold=True, color=accent)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(4)
        set_run(p2.add_run(label), 8.1, color=BODY)
    add_para(doc, "", after=3)
    return table


def add_feature(doc, label, text, after=3):
    return add_rich(
        doc,
        [
            ("⬢  ", {"bold": True, "color": MOSS, "size": 9.5}),
            (label + "  ", {"bold": True, "color": INK}),
            (text, {"color": BODY}),
        ],
        after=after,
        keep=True,
    )


def add_dark_banner(doc, label, text, after=6):
    table = doc.add_table(rows=1, cols=1)
    base.set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    base.set_cell_shading(cell, DARK)
    base.set_cell_border(cell, DARK, size=8)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    set_run(p.add_run(label.upper() + "  "), 8.5, bold=True, color=HONEY_P)
    set_run(p.add_run(text), 11, bold=True, color=WHITE)
    add_para(doc, "", after=after)
    return table


def add_revenue_strip(doc):
    entries = [
        ("MODEL", "one-time integration"),
        ("DEVICE", "embedded or gateway licence"),
        ("FLEET", "per-active-device software"),
        ("PROGRAM", "enterprise fee + optional value share"),
    ]
    widths = [2340, 2340, 2340, 2340]
    table = doc.add_table(rows=1, cols=4)
    base.set_table_geometry(table, widths)
    for i, (label, body) in enumerate(entries):
        cell = table.cell(0, i)
        base.set_cell_shading(cell, CREAM if i % 2 == 0 else PAPER)
        base.set_cell_border(cell, LINE, size=5)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(label), 7.8, bold=True, color=HONEY_D)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(3)
        set_run(p2.add_run(body), 8.5, bold=True, color=INK)
    add_para(doc, "", after=3)
    return table


def add_numbered_step(doc, text, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), str(doc._kwh_decimal_num_id))
    numPr.extend([ilvl, numId])
    pPr.append(numPr)
    set_run(p.add_run(text), 9.6, color=BODY)
    return p


def add_hyperlink(paragraph, text, url, size=7.5):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rPr.append(rFonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), HONEY_D)
    rPr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def build():
    make_hex_band()
    doc = Document()
    configure(doc)

    add_hex_band(doc)
    add_kicker(doc, "Battery OEM partner brief · India · July 2026")
    add_title(doc, "The recurring-services layer for the batteries you ship.", size=27, after=6)
    add_para(
        doc,
        "kWh Electric connects compatible battery systems to approved DISCOM programs, authorised trading providers and flexibility services through one policy-bound edge runtime and one fleet API.",
        size=12,
        color=BODY,
        after=8,
    )
    add_dark_banner(
        doc,
        "The wedge",
        "Trading applications can match energy. kWh Electric makes the physical battery observable, controllable, permissioned and verifiable enough to deliver the schedule.",
        after=3,
    )

    add_kicker(doc, "01 · The participation layer", after=2)
    add_title(doc, "What kWh Electric is", size=17, after=4)
    add_para(
        doc,
        "An OEM-neutral control layer that lives beside—or inside—the battery controller. It normalises device capabilities, protects the owner reserve and warranty envelope, executes approved schedules locally, and reconciles device telemetry with the regulated boundary meter.",
        size=10.2,
        after=4,
    )
    add_feature(doc, "Interoperable", "Reusable profiles translate CAN, Modbus and inverter/BMS data into one consistent asset model.")
    add_feature(doc, "Permissioned", "Every service receives purpose-limited access; owner opt-out, BMS safety, OEM limits and DISCOM authority always outrank a market command.")
    add_feature(doc, "Verifiable", "Availability, renewable-charge provenance, command acknowledgement, exceptions and meter variance create an auditable settlement trail.", after=4)

    add_kicker(doc, "02 · Two ways to ship", after=2)
    add_title(doc, "Embedded runtime or gateway—one control contract", size=16, after=4)
    add_compare(
        doc,
        ["Embedded with the OEM", "Shipped as a gateway"],
        [
            "The kWh runtime is integrated into a supported controller or connectivity module. Best long-term route for volume and the cleanest OEM customer experience.",
            "A compact edge device connects to the exact approved BMS/inverter interface. Best first route for pilots, mixed fleets and faster field validation.",
        ],
    )
    add_rich(
        doc,
        [
            ("Boundary: ", {"bold": True, "color": HONEY_D}),
            ("kWh does not replace the OEM app, BMS, DISCOM meter, bill or trading application. It makes those systems work together safely.", {"color": BODY}),
        ],
        size=9.4,
        after=0,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_hex_band(doc, page_break_before=True)
    add_kicker(doc, "03 · The value stack", after=2)
    add_title(doc, "Three services. One connected battery.", size=24, after=5)
    add_metric_strip(
        doc,
        [
            ("BESS", "renewable-charged storage is eligible in Delhi P2P rules", HONEY_P, HONEY_D),
            ("11", "authorised trading applications listed in the PVVNL IES pilot", MOSS_P, MOSS),
            ("40 lakh+", "PM Surya Ghar beneficiary households by May 2026", CREAM, HONEY_D),
        ],
    )

    add_feature(doc, "1 · P2P market participation", "Offer renewable-eligible battery energy through an authorised provider while the DISCOM retains meter verification and bill settlement. Owner: protected backup plus incremental value. OEM: a differentiated connected service.", after=3)
    add_feature(doc, "2 · Demand flexibility", "Aggregate safe residual capacity for peak support, local flexibility and future DISCOM programs. Owner: program payments where available. OEM: recurring fleet and orchestration revenue.", after=3)
    add_feature(doc, "3 · Connected asset services", "Use the same telemetry, permissions and policy engine for fleet health, warranty-aware control, installer support and portfolio reporting. OEM: a durable software relationship beyond the hardware sale.", after=5)

    add_dark_banner(
        doc,
        "Commercial principle",
        "Transaction fees are not the business. kWh earns by enabling and operating the asset: model integration, device software, fleet orchestration and enterprise program assurance.",
        after=3,
    )
    add_revenue_strip(doc)

    add_kicker(doc, "04 · Proposed first engagement", after=2)
    add_title(doc, "A 90-day Delhi battery-participation pilot", size=16, after=3)
    add_numbered_step(doc, "Days 0–30 — integrate one exact OEM SKU; validate read/write control, export mode, owner reserve, warranty constraints and renewable-energy accounting on three lab systems.")
    add_numbered_step(doc, "Days 31–60 — connect one authorised trading provider and a TPDDL or BRPL pilot route; deploy an initial 25–50-site cohort only after the lab acceptance gate.")
    add_numbered_step(doc, "Days 61–90 — submit availability, execute accepted schedules, reconcile the DISCOM meter and publish a joint result pack covering delivery, safety, owner value and degradation.", after=3)
    add_rich(
        doc,
        [
            ("Success gate: ", {"bold": True, "color": HONEY_D}),
            ("≥95% command acknowledgement · <2% gateway-to-meter variance · positive owner value after losses and degradation · zero reserve, warranty or thermal violations.", {"bold": True, "color": INK}),
        ],
        size=9.2,
        after=4,
    )

    add_rich(
        doc,
        [
            ("Seeking: ", {"bold": True, "color": HONEY_D}),
            ("one battery OEM, one authorised market partner and one Delhi DISCOM sponsor.  ", {"bold": True, "color": INK}),
            ("arham@kwhelectric.io  ·  kwhelectric.io", {"bold": True, "color": HONEY_D}),
        ],
        size=9.4,
        after=3,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run("Evidence: "), 7.3, bold=True, color=MUTED)
    links = [
        ("DERC P2P Guidelines 2024", base.SOURCES[1][1]),
        ("DERC Order 02/2026", base.SOURCES[2][1]),
        ("PVVNL IES pilot", base.SOURCES[3][1]),
        ("MNRE/PIB", base.SOURCES[5][1]),
    ]
    for i, (label, url) in enumerate(links):
        add_hyperlink(p, label, url, size=7.3)
        if i < len(links) - 1:
            set_run(p.add_run("  ·  "), 7.3, color=MUTED)

    EXPORTS.mkdir(parents=True, exist_ok=True)
    path = EXPORTS / "kWh-Electric-Battery-OEM-Partner-Brief-v2.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build())
